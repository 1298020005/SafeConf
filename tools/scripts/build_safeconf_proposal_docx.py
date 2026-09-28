#!/usr/bin/env python3
"""Render this proposal's limited Markdown syntax into an editable Word file."""
import argparse
import re
from pathlib import Path
from urllib.parse import quote

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

REPO = Path(__file__).resolve().parents[2]
TOKEN = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')


def add_inline(paragraph, text, source):
    for token in TOKEN.split(text):
        if not token:
            continue
        link = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token)
        if link:
            label, target = link.groups()
            if not target.startswith(('https://', 'http://')):
                path = (source.parent / target).resolve()
                target = 'https://github.com/1298020005/SafeConf/blob/exp/e220-reviewer-closure-20260921/' + quote(str(path.relative_to(REPO)))
            hyperlink = OxmlElement('w:hyperlink')
            hyperlink.set(qn('r:id'), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
            run = OxmlElement('w:r')
            props = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '216C97'); props.append(color)
            run.append(props)
            node = OxmlElement('w:t'); node.text = label; run.append(node)
            hyperlink.append(run); paragraph._p.append(hyperlink)
        elif token.startswith('**'):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith('`'):
            run = paragraph.add_run(token[1:-1]); run.font.name = 'DejaVu Sans Mono'; run.font.size = Pt(9)
        else:
            paragraph.add_run(token)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    source = args.source.resolve()
    lines = source.read_text().splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.0)
    section.left_margin = section.right_margin = Cm(2.1)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number']:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    normal = doc.styles['Normal']; normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name in ['Heading 1', 'Heading 2', 'Heading 3']:
        doc.styles[name].font.color.rgb = RGBColor.from_string('174861')
    header = section.header.paragraphs[0]
    header.text = 'SafeConf｜证据驱动的方法研究与开题进展｜2026-09-28'
    header.style = doc.styles['Normal']
    for run in header.runs:
        run.font.size = Pt(8)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    i, code = 0, False
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith('```'):
            code = not code; i += 1; continue
        if not line:
            i += 1; continue
        if code:
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
            run = p.add_run(line); run.font.name = 'DejaVu Sans Mono'; run.font.size = Pt(9)
            i += 1; continue
        img = re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)', line)
        if img:
            doc.add_picture(str(source.parent / img[2]), width=Cm(16.6))
            p = doc.add_paragraph(img[1]); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue
        if line.startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                candidate = lines[i].strip()
                if not re.fullmatch(r'[|:\s-]+', candidate):
                    table_lines.append([c.strip() for c in candidate.strip('|').split('|')])
                i += 1
            columns = len(table_lines[0])
            assert all(len(r) == columns for r in table_lines)
            table = doc.add_table(rows=0, cols=columns); table.style = 'Light Shading Accent 1'
            for row_index, values in enumerate(table_lines):
                cells = table.add_row().cells
                for cell, value in zip(cells, values):
                    p = cell.paragraphs[0]; add_inline(p, value, source)
                    for run in p.runs:
                        run.font.size = Pt(9)
                        if row_index == 0: run.bold = True
                if row_index == 0:
                    repeat = OxmlElement('w:tblHeader'); table.rows[-1]._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph()
            continue
        heading = re.match(r'^(#{1,4})\s+(.*)', line)
        if heading:
            level = len(heading[1])
            p = doc.add_paragraph(style='Title' if level == 1 else f'Heading {level - 1}')
            add_inline(p, heading[2], source)
        elif line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet'); add_inline(p, line[2:], source)
        elif re.match(r'^\d+\. ', line):
            p = doc.add_paragraph(style='List Number'); add_inline(p, re.sub(r'^\d+\. ', '', line), source)
        else:
            p = doc.add_paragraph(); add_inline(p, line, source)
        i += 1
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output)


if __name__ == '__main__':
    main()
